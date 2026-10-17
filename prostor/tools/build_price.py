#!/usr/bin/env python3
"""Собирает assets/price.js из прайс-листа питомника (xlsx).

    python3 tools/build_price.py docs/prostor-price.xlsx --date 12.08.2025

Структура листа: группы (уровень 1) и роды (уровень 2) — строки с заливкой
без цены; позиции — строки с ценой: A — номенклатура, M — характеристика,
N — стоимость.
"""
import argparse, json, os
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

ap = argparse.ArgumentParser()
ap.add_argument("xlsx")
ap.add_argument("--date", help="дата прайса, по умолчанию — дата изменения файла")
ap.add_argument("--out", default=os.path.join(ROOT, "assets", "price.js"))
args = ap.parse_args()

wb = openpyxl.load_workbook(args.xlsx, data_only=True)
ws = wb.worksheets[0]

groups, group, genus = [], None, None
for r in range(1, ws.max_row + 1):
    name, char, price = ws.cell(r, 1).value, ws.cell(r, 13).value, ws.cell(r, 14).value
    if name is None:
        continue
    name = " ".join(str(name).split())
    level = ws.row_dimensions[r].outlineLevel
    if price is None:
        if level == 1:
            group = [name, []]; groups.append(group); genus = None
        elif level == 2 and group:
            genus = [name, []]; group[1].append(genus)
        continue
    if not isinstance(price, (int, float)) or genus is None:
        continue
    price = round(price, 2)
    genus[1].append([name, " ".join(str(char or "").split()), int(price) if price == int(price) else price])

date = args.date or (wb.properties.modified.strftime("%d.%m.%Y") if wb.properties.modified else "")
data = {"date": date, "groups": groups}
total = sum(len(g[1]) for gr in groups for g in gr[1])

with open(args.out, "w", encoding="utf-8") as f:
    f.write("/* Сгенерировано tools/build_price.py — вручную не редактировать */\n")
    f.write("window.PRICE=" + json.dumps(data, ensure_ascii=False, separators=(",", ":")) + ";\n")
print(f"{args.out}: {len(groups)} групп, {total} позиций, прайс от {date}")
